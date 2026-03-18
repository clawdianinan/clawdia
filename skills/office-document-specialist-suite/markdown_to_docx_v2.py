#!/usr/bin/env python3
"""
Convert markdown to professionally formatted DOCX document.
Simplified version.
"""

import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from docx.enum.style import WD_STYLE_TYPE


def clean_markdown(text):
    """Remove markdown formatting while preserving structure."""
    if not text:
        return text
    
    # Remove bold/italic markers
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'__(.*?)__', r'\1', text)
    text = re.sub(r'_(.*?)_', r'\1', text)
    
    # Remove inline code markers
    text = re.sub(r'`(.*?)`', r'\1', text)
    
    # Remove markdown links but keep text
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    
    return text.strip()


def apply_document_styles(doc):
    """Apply professional styling to the document."""
    
    # Configure Normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.15
    
    # Configure Heading 1
    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Calibri'
    h1_style.font.size = Pt(18)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)  # Dark blue
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(8)
    h1_style.paragraph_format.line_spacing = 1.2
    
    # Configure Heading 2
    h2_style = doc.styles['Heading 2']
    h2_style.font.name = 'Calibri'
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0x2F, 0x55, 0x8C)  # Medium blue
    h2_style.paragraph_format.space_before = Pt(10)
    h2_style.paragraph_format.space_after = Pt(6)
    h2_style.paragraph_format.line_spacing = 1.2
    
    # Configure Heading 3
    h3_style = doc.styles['Heading 3']
    h3_style.font.name = 'Calibri'
    h3_style.font.size = Pt(12)
    h3_style.font.bold = True
    h3_style.font.color.rgb = RGBColor(0x4F, 0x6B, 0xA6)  # Light blue
    h3_style.paragraph_format.space_before = Pt(8)
    h3_style.paragraph_format.space_after = Pt(4)
    h3_style.paragraph_format.line_spacing = 1.15
    
    # Create Title style
    if 'Title' not in doc.styles:
        title_style = doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
    else:
        title_style = doc.styles['Title']
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)  # Black
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Configure page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)


def convert_markdown_to_docx(markdown_path, docx_path):
    """Convert markdown file to professionally formatted DOCX."""
    
    print(f"Reading markdown file: {markdown_path}")
    with open(markdown_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split into lines for processing
    lines = content.split('\n')
    
    # Create document
    doc = Document()
    apply_document_styles(doc)
    
    # Process each line
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            i += 1
            continue
        
        # Check for headings
        if line.startswith('# '):
            # Main title
            title_text = clean_markdown(line[2:])
            p = doc.add_paragraph(title_text, style='Title')
            i += 1
            
        elif line.startswith('## '):
            # Heading 1
            heading_text = clean_markdown(line[3:])
            doc.add_paragraph(heading_text, style='Heading 1')
            i += 1
            
        elif line.startswith('### '):
            # Heading 2
            heading_text = clean_markdown(line[4:])
            doc.add_paragraph(heading_text, style='Heading 2')
            i += 1
            
        elif line.startswith('#### '):
            # Heading 3
            heading_text = clean_markdown(line[5:])
            doc.add_paragraph(heading_text, style='Heading 3')
            i += 1
            
        elif line.startswith('- ') or line.startswith('* '):
            # List item
            list_text = clean_markdown(line[2:])
            p = doc.add_paragraph(list_text, style='Normal')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            i += 1
            
        elif re.match(r'^\d+\.\s', line):
            # Numbered list item
            list_text = clean_markdown(re.sub(r'^\d+\.\s', '', line))
            p = doc.add_paragraph(list_text, style='Normal')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            i += 1
            
        elif line.startswith('> '):
            # Blockquote
            quote_text = clean_markdown(line[2:])
            p = doc.add_paragraph(quote_text, style='Normal')
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            
        elif line.startswith('---'):
            # Horizontal rule - skip
            i += 1
            
        else:
            # Regular paragraph - collect consecutive lines
            paragraph_lines = []
            while i < len(lines) and lines[i].strip() and not any([
                lines[i].strip().startswith('#'),
                lines[i].strip().startswith('- '),
                lines[i].strip().startswith('* '),
                re.match(r'^\d+\.\s', lines[i].strip()),
                lines[i].strip().startswith('> '),
                lines[i].strip().startswith('---'),
            ]):
                paragraph_lines.append(clean_markdown(lines[i]))
                i += 1
            
            if paragraph_lines:
                paragraph_text = ' '.join(paragraph_lines)
                doc.add_paragraph(paragraph_text, style='Normal')
            else:
                i += 1
    
    # Save the document
    print(f"Saving DOCX file: {docx_path}")
    doc.save(docx_path)
    print(f"Conversion complete! File saved to: {docx_path}")


if __name__ == '__main__':
    import argparse
    
    parser = argparse.ArgumentParser(description='Convert markdown to professionally formatted DOCX')
    parser.add_argument('input', help='Input markdown file path')
    parser.add_argument('output', help='Output DOCX file path')
    
    args = parser.parse_args()
    
    convert_markdown_to_docx(args.input, args.output)