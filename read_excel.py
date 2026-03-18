#!/usr/bin/env python3
import sys
import os

# Try to read Excel file without pandas
excel_file = "IIH_Roles_and_Grades_20260224_FINAL_CORRECTED_20260224_1839.xlsx"

if not os.path.exists(excel_file):
    print(f"File not found: {excel_file}")
    sys.exit(1)

print(f"File exists: {excel_file}")
print(f"File size: {os.path.getsize(excel_file)} bytes")

# Try to use openpyxl if available
try:
    import openpyxl
    print("openpyxl is available")
    
    # Load the workbook
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    
    print(f"\nSheets in workbook: {wb.sheetnames}")
    
    # Check for summary sheet
    summary_sheets = [sheet for sheet in wb.sheetnames if 'summary' in sheet.lower()]
    if summary_sheets:
        print(f"\nSummary sheet found: {summary_sheets[0]}")
        ws = wb[summary_sheets[0]]
        
        print(f"\nReading first 20 rows of summary sheet:")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
            print(f"Row {i}: {row}")
    else:
        print("\nNo summary sheet found. Checking all sheets:")
        for sheet_name in wb.sheetnames[:3]:  # Check first 3 sheets
            ws = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} ---")
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), 1):
                print(f"Row {i}: {row}")
                
except ImportError:
    print("openpyxl not available")
    print("Trying to install openpyxl...")
    import subprocess
    result = subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "--user"], 
                          capture_output=True, text=True)
    print(f"Install result: {result.returncode}")
    if result.returncode == 0:
        print("openpyxl installed successfully")
        import openpyxl
        # Retry loading
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
    else:
        print(f"Failed to install openpyxl: {result.stderr}")